#!/usr/bin/env python3
"""Generate the mail-merge test fixtures (run: python3 -I tests/fixtures/mailmerge/make_fixtures.py).

Hand-written OOXML / ODF parts zipped with zipfile, plus one workbook written by openpyxl when it is
installed, so the parser is exercised on both minimal and real-producer files.
"""
import io
import os
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))


def out(name):
    return os.path.join(HERE, name)


# ------------------------------------------------------------------------------------------------
# people.xlsx: shared strings (plain + rich runs + phonetic), inline strings, str/b/e cells, dates,
# times, date-times, number formats, sparse cells, a leading blank row, multiple sheets in a
# workbook order that differs from the rels order, a hidden sheet, a chart-less 1904 workbook flag off.
# ------------------------------------------------------------------------------------------------
CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
</Types>"""

ROOT_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>"""

WORKBOOK = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<workbookPr date1904="false"/>
<sheets>
<sheet name="Recipients" sheetId="1" r:id="rId3"/>
<sheet name="Second &amp; More" sheetId="2" r:id="rId1"/>
<sheet name="Secret" sheetId="3" state="hidden" r:id="rId2"/>
</sheets>
</workbook>"""

WB_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet2.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="/xl/worksheets/sheet3.xml"/>
<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
<Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>
<Relationship Id="rId5" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
</Relationships>"""

SHARED = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="9" uniqueCount="9">
<si><t>First Name</t></si>
<si><t>Last Name</t></si>
<si><t>City</t></si>
<si><t>Joined</t></si>
<si><r><rPr><b/></rPr><t>Ada</t></r><r><t xml:space="preserve"> Grace</t></r></si>
<si><t>Lovelace</t><rPh sb="0" eb="1"><t>IGNORED</t></rPh></si>
<si><t>London &amp; Paris</t></si>
<si><t xml:space="preserve">  spaced  </t></si>
<si><t>Line1_x000A_Line2</t></si>
</sst>"""

STYLES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<numFmts count="3">
<numFmt numFmtId="164" formatCode="dd/mm/yyyy"/>
<numFmt numFmtId="165" formatCode="&quot;$&quot;#,##0.00"/>
<numFmt numFmtId="166" formatCode="[$-409]h:mm:ss AM/PM"/>
</numFmts>
<cellXfs count="8">
<xf numFmtId="0"/>
<xf numFmtId="14"/>
<xf numFmtId="164"/>
<xf numFmtId="165"/>
<xf numFmtId="22"/>
<xf numFmtId="166"/>
<xf numFmtId="10"/>
<xf numFmtId="3"/>
</cellXfs>
</styleSheet>"""

# Sheet 1 "Recipients": blank first row, header on row 2, sparse cells, cells without r attributes.
SHEET1 = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<dimension ref="A2:I6"/>
<sheetData>
<row r="1"><c r="B1" s="0"/></row>
<row r="2">
  <c r="A2" t="s"><v>0</v></c><c r="B2" t="s"><v>1</v></c><c r="C2" t="s"><v>2</v></c><c r="D2" t="s"><v>3</v></c>
  <c r="E2" t="inlineStr"><is><t>Balance</t></is></c><c r="F2" t="inlineStr"><is><r><t>Act</t></r><r><t>ive</t></r></is></c>
  <c r="G2" t="str"><v>Note</v></c><c r="H2" t="inlineStr"><is><t>First Name</t></is></c><c r="I2"/>
</row>
<row r="3">
  <c r="A3" t="s"><v>4</v></c><c r="B3" t="s"><v>5</v></c><c r="C3" t="s"><v>6</v></c><c r="D3" s="1"><v>45306</v></c>
  <c r="E3" s="3"><v>1234.5</v></c><c r="F3" t="b"><v>1</v></c><c r="G3" t="str"><f>CONCAT("a","b")</f><v>ab</v></c><c r="H3" t="s"><v>8</v></c>
</row>
<row r="5">
  <c r="A5" t="inlineStr"><is><t>Bob</t></is></c><c r="D5" s="2"><v>36526</v></c><c r="E5"><v>0.30000000000000004</v></c><c r="F5" t="b"><v>0</v></c>
  <c r="G5" t="e"><v>#N/A</v></c>
</row>
<row r="6">
  <c t="s"><v>7</v></c><c><v>42</v></c><c t="s"><v>6</v></c><c s="4"><v>45306.75</v></c><c s="6"><v>0.125</v></c><c s="7"><v>1234567</v></c>
</row>
<row r="9"><c r="A9" t="inlineStr"><is><t></t></is></c></row>
</sheetData>
</worksheet>"""

SHEET2 = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<x:worksheet xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<x:sheetData>
<x:row r="1"><x:c r="A1" t="inlineStr"><x:is><x:t>Code</x:t></x:is></x:c><x:c r="C1" t="inlineStr"><x:is><x:t>Time</x:t></x:is></x:c><x:c r="D1" t="inlineStr"><x:is><x:t>Code</x:t></x:is></x:c></x:row>
<x:row r="2"><x:c r="A2"><x:v>7</x:v></x:c><x:c r="B2"><x:v>8</x:v></x:c><x:c r="C2" s="5"><x:v>0.5625</x:v></x:c><x:c r="D2" t="d"><x:v>2024-03-05T00:00:00</x:v></x:c></x:row>
<x:row r="3"><x:c r="AA3" t="inlineStr"><x:is><x:t>far</x:t></x:is></x:c></x:row>
</x:sheetData>
</x:worksheet>"""

SHEET3 = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>
<row r="1"><c r="A1" t="inlineStr"><is><t>Key</t></is></c></row>
<row r="2"><c r="A2" t="inlineStr"><is><t>hidden value</t></is></c></row>
</sheetData></worksheet>"""


def write_xlsx(path, compression=zipfile.ZIP_DEFLATED, stream=False):
    parts = [
        ('[Content_Types].xml', CONTENT_TYPES), ('_rels/.rels', ROOT_RELS), ('xl/workbook.xml', WORKBOOK),
        ('xl/_rels/workbook.xml.rels', WB_RELS), ('xl/sharedStrings.xml', SHARED), ('xl/styles.xml', STYLES),
        ('xl/worksheets/sheet1.xml', SHEET1), ('xl/worksheets/sheet2.xml', SHEET2), ('xl/worksheets/sheet3.xml', SHEET3),
    ]
    if stream:
        # a non-seekable target makes zipfile write data descriptors (general purpose flag bit 3)
        class Unseekable(io.RawIOBase):
            def __init__(self):
                self.buf = bytearray()

            def writable(self):
                return True

            def write(self, b):
                self.buf.extend(b)
                return len(b)

            def seekable(self):
                return False

        target = Unseekable()
        with zipfile.ZipFile(target, 'w', compression) as z:
            for name, data in parts:
                with z.open(name, 'w') as f:
                    f.write(data.encode('utf-8'))
        with open(path, 'wb') as f:
            f.write(bytes(target.buf))
        return
    with zipfile.ZipFile(path, 'w', compression) as z:
        for name, data in parts:
            z.writestr(name, data)


write_xlsx(out('people.xlsx'))
write_xlsx(out('people-stored.xlsx'), zipfile.ZIP_STORED)
write_xlsx(out('people-stream.xlsx'), zipfile.ZIP_DEFLATED, stream=True)

# ZIP64 variant: lower zipfile's limits so it writes ZIP64 end records and ZIP64 extra fields.
_saved = (zipfile.ZIP64_LIMIT, zipfile.ZIP_FILECOUNT_LIMIT)
zipfile.ZIP64_LIMIT = 0
zipfile.ZIP_FILECOUNT_LIMIT = 0
try:
    write_xlsx(out('people-zip64.xlsx'))
finally:
    zipfile.ZIP64_LIMIT, zipfile.ZIP_FILECOUNT_LIMIT = _saved

# ------------------------------------------------------------------------------------------------
# openpyxl-written workbook (real producer: shared strings, styles, dates as datetime)
# ------------------------------------------------------------------------------------------------
try:
    import datetime
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Members'
    ws.append(['Name', 'Email', 'Birthday', 'Fee', 'Member', 'Zip'])
    ws.append(['Zoë Ångström', 'zoe@example.com', datetime.date(1990, 7, 4), 12.5, True, '02134'])
    ws.append(['Li Wei', None, datetime.datetime(1985, 12, 31, 14, 30), 7, False, '10001'])
    ws.append([None, None, None, None, None, None])
    ws.append(['Omar', 'omar@example.com', None, None, None, None])
    ws['D2'].number_format = '"$"#,##0.00'
    ws2 = wb.create_sheet('Empty')
    wb.save(out('openpyxl.xlsx'))
except ImportError:
    pass

# ------------------------------------------------------------------------------------------------
# people.ods
# ------------------------------------------------------------------------------------------------
ODS_CONTENT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
 xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:dc="http://purl.org/dc/elements/1.1/" office:version="1.2">
<office:automatic-styles>
<style:style style:name="ta1" style:family="table"><style:table-properties table:display="true"/></style:style>
<style:style style:name="ta2" style:family="table"><style:table-properties table:display="false"/></style:style>
</office:automatic-styles>
<office:body><office:spreadsheet>
<table:table table:name="Guests" table:style-name="ta1">
<table:table-column table:number-columns-repeated="1024"/>
<table:table-row table:number-rows-repeated="2"><table:table-cell table:number-columns-repeated="1024"/></table:table-row>
<table:table-row>
<table:table-cell office:value-type="string"><text:p>Name</text:p></table:table-cell>
<table:table-cell office:value-type="string"><text:p>Date</text:p></table:table-cell>
<table:table-cell office:value-type="string"><text:p>Amount</text:p></table:table-cell>
<table:table-cell/>
<table:table-cell office:value-type="string"><text:p>Flag</text:p></table:table-cell>
<table:table-cell office:value-type="string"><text:p>Time</text:p></table:table-cell>
<table:table-cell table:number-columns-repeated="1018"/>
</table:table-row>
<table:table-row>
<table:table-cell office:value-type="string"><text:p>Ann<text:s text:c="2"/>Lee</text:p><office:annotation><dc:creator>x</dc:creator><text:p>NOTE</text:p></office:annotation></table:table-cell>
<table:table-cell office:value-type="date" office:date-value="2024-01-15"><text:p>01/15/24</text:p></table:table-cell>
<table:table-cell office:value-type="currency" office:currency="USD" office:value="1234.5"><text:p>$1,234.50</text:p></table:table-cell>
<table:table-cell office:value-type="string"><text:p>extra</text:p></table:table-cell>
<table:table-cell office:value-type="boolean" office:boolean-value="true"><text:p>TRUE</text:p></table:table-cell>
<table:table-cell office:value-type="time" office:time-value="PT13H30M00S"><text:p>01:30 PM</text:p></table:table-cell>
</table:table-row>
<table:table-row table:number-rows-repeated="2">
<table:table-cell table:number-columns-repeated="2" office:value-type="string"><text:p>dup</text:p></table:table-cell>
<table:table-cell office:value-type="float" office:value="3"/>
</table:table-row>
<table:table-row>
<table:table-cell office:value-type="string"><text:p>Two</text:p><text:p>Lines</text:p></table:table-cell>
<table:table-cell office:value-type="date" office:date-value="2024-02-29T08:05:00"><text:p>x</text:p></table:table-cell>
<table:table-cell office:value-type="percentage" office:value="0.25"><text:p>25%</text:p></table:table-cell>
</table:table-row>
<table:table-row table:number-rows-repeated="1048570"><table:table-cell table:number-columns-repeated="1024"/></table:table-row>
</table:table>
<table:table table:name="Hidden Sheet" table:style-name="ta2">
<table:table-row><table:table-cell office:value-type="string"><text:p>H</text:p></table:table-cell></table:table-row>
<table:table-row><table:table-cell office:value-type="float" office:value="1.5"><text:p>1.5</text:p></table:table-cell></table:table-row>
</table:table>
</office:spreadsheet></office:body></office:document-content>"""

MANIFEST = """<?xml version="1.0" encoding="UTF-8"?>
<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2">
<manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>
<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>
</manifest:manifest>"""

with zipfile.ZipFile(out('people.ods'), 'w') as z:
    z.writestr(zipfile.ZipInfo('mimetype'), 'application/vnd.oasis.opendocument.spreadsheet', compress_type=zipfile.ZIP_STORED)
    z.writestr('content.xml', ODS_CONTENT, compress_type=zipfile.ZIP_DEFLATED)
    z.writestr('META-INF/manifest.xml', MANIFEST, compress_type=zipfile.ZIP_DEFLATED)

# ------------------------------------------------------------------------------------------------
# legacy .xls (OLE2 compound file signature) and text files
# ------------------------------------------------------------------------------------------------
with open(out('legacy.xls'), 'wb') as f:
    f.write(bytes([0xD0, 0xCF, 0x11, 0xE0, 0xA1, 0xB1, 0x1A, 0xE1]) + bytes(504))

with open(out('cp1252.csv'), 'wb') as f:
    f.write('Name;City;Price\r\nJosé;Zürich;€5\r\n"Smith; Jr.";Köln;“quoted”\r\n'.encode('cp1252'))

with open(out('bom-utf8.csv'), 'wb') as f:
    f.write('﻿Name,Note\n"Ann","multi\nline"\nBob,ünïcödé\n'.encode('utf-8'))

with open(out('utf16.txt'), 'wb') as f:
    f.write('﻿A\tB\r\n1\t2\r\n'.encode('utf-16-le'))

with open(out('tabs.tsv'), 'wb') as f:
    f.write(b'Name\tCity, State\nAnn\tBoston, MA\n')

print('fixtures written to', HERE)
